import os

from docx import Document

from config.settings import UPLOAD_FOLDER



def extract_examinee_info(doc_path):
    doc = Document(doc_path)
    examinee_info = {}

    for table in doc.tables:
        first_cell_text = table.cell(0, 0).text.strip()
        print(f"Checking table with first cell: '{first_cell_text}'")  # Debug statement

        if "Examinee Name" in first_cell_text or "Date of Testing" in first_cell_text:  # Adjusted to include both tables
            for row in table.rows:
                cell_texts = [cell.text.strip() for cell in row.cells]
                print(cell_texts)
                print(f"Row content: {cell_texts}")  # Print each row's content

                # Extract pairs based on the expected layout
                for i in range(0, len(row.cells), 2):  # Handles pairs (step by 2)
                    if i + 1 < len(row.cells):  # Check if there's a corresponding value cell
                        key = row.cells[i].text.strip()
                        value = row.cells[i+1].text.strip()
                        if key:  # Only add to dict if key exists

                            examinee_info[key] = value
                            if "Date of Testing" in key:
                                print('fart')
                                key = row.cells[3].text.strip()
                                value = row.cells[4].text.strip()
                                print("Key =", key, " Value =", value)
                                examinee_info[key] = value
                                break
                            print(f"Extracted: {key} - {value}")  # Debug statement


    replacements = {}
    replacements['[Full Name]'] = examinee_info['Examinee Name']
    replacements['[First Name]'] = examinee_info['Examinee Name'].split(" ")[0]
    replacements['[Exact Age]'] = examinee_info['Age at Testing']
    replacements['[Dates of Evaluation]'] = examinee_info['Date of Testing']

    return replacements

# Usage
doc_path = os.path.join(UPLOAD_FOLDER, 'wais.docx')
info = extract_examinee_info(doc_path)
print("Extracted Information:", info)
